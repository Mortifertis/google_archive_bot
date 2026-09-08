"""Tests for building DOCX documents from Telegram posts."""

import base64
from datetime import datetime, timezone
from unittest.mock import patch

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from app.document_builder import (MAX_TITLE_LENGTH, build_document_title,
                                  build_post_docx)
from app.telegram_media import DownloadedPhoto
from app.telegram_parser import ForwardedPost

PNG_DATA = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk"
    "/wcAAusB9Wl2nF0AAAAASUVORK5CYII="
)
FOOTER_TEXT = "Сохранено через Telegram Archive Bot"


def _photo(
    unique_id: str = "photo",
    data: bytes = PNG_DATA,
) -> DownloadedPhoto:
    return DownloadedPhoto(data, 1, 1, unique_id)


def _post(**changes: object) -> ForwardedPost:
    values = {
        "source_chat_id": -100123,
        "source_chat_title": "Some Channel",
        "source_chat_username": "somechannel",
        "source_message_id": 12345,
        "source_date": datetime(2026, 9, 7, 18, 40, tzinfo=timezone.utc),
        "source_url": "https://t.me/somechannel/12345",
        "text": "\n  Первая   содержательная строка  \nВторой абзац",
        "caption": None,
        "media_group_id": None,
        "photo_count": 0,
        "is_channel_post": True,
    }
    values.update(changes)
    return ForwardedPost(**values)


def _paragraph(document: Document, text: str):
    return next(item for item in document.paragraphs if item.text == text)


def test_title_contains_date_channel_and_first_meaningful_line() -> None:
    title = build_document_title(_post())

    assert title == (
        "2026-09-07 — Some Channel — Первая содержательная строка"
    )


def test_drive_title_is_limited_without_changing_its_logic() -> None:
    title = build_document_title(_post(text="x" * 300))

    assert len(title) == MAX_TITLE_LENGTH
    assert title.endswith("…")


def test_document_uses_editorial_header_and_compact_metadata() -> None:
    document = Document(
        build_post_docx(_post(text="Короткий заголовок\n\nBody"))
    )
    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    assert paragraphs[:4] == [
        "SOME CHANNEL",
        "Короткий заголовок",
        "07.09.2026 18:40 · Telegram",
        "https://t.me/somechannel/12345",
    ]
    assert "Короткий заголовок" not in paragraphs[2:]
    assert "Body" in paragraphs
    assert not any(text.startswith("Источник:") for text in paragraphs)
    assert "Изображения" not in paragraphs
    assert build_document_title(_post()) not in paragraphs


def test_missing_date_and_url_adds_only_telegram_metadata() -> None:
    document = Document(
        build_post_docx(_post(source_date=None, source_url=None))
    )
    paragraphs = [paragraph.text for paragraph in document.paragraphs]

    assert "Telegram" in paragraphs
    assert "недоступен" not in paragraphs
    assert "недоступно" not in paragraphs


def test_normal_and_body_typography() -> None:
    document = Document(
        build_post_docx(_post(text="Заголовок\n\nОбычная проза"))
    )
    normal = document.styles["Normal"]
    body = _paragraph(document, "Обычная проза")

    assert normal.font.name == "Noto Sans"
    assert normal.font.size == Pt(12.5)
    assert body.runs[0].font.name == "Noto Sans"
    assert body.runs[0].font.size == Pt(12.5)
    assert body.paragraph_format.first_line_indent.cm == pytest.approx(
        1.25, abs=0.001
    )
    assert body.paragraph_format.space_after == Pt(4)
    assert body.paragraph_format.line_spacing == pytest.approx(1.2)
    assert body.alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_blank_lines_create_only_logical_body_paragraphs() -> None:
    content = (
        "Абзац один, достаточно длинный для отсутствия заголовка "
        + "x" * 80
        + "\n\n\nАбзац два.\n\n\n\nАбзац три."
    )
    document = Document(build_post_docx(_post(text=content)))
    body = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.runs
        and paragraph.runs[0].font.size == Pt(12.5)
    ]

    assert body == [
        "Абзац один, достаточно длинный для отсутствия заголовка "
        + "x" * 80,
        "Абзац два.",
        "Абзац три.",
    ]


@pytest.mark.parametrize(
    "text",
    [
        "• пункт",
        "● пункт",
        "◦ пункт",
        "- пункт",
        "— пункт",
        "* пункт",
        "1. пункт",
        "2) пункт",
    ],
)
def test_structural_body_lines_have_no_first_line_indent(text: str) -> None:
    long_first_line = "Длинное вступление " + "x" * 120
    document = Document(
        build_post_docx(_post(text=f"{long_first_line}\n\n{text}"))
    )

    assert _paragraph(
        document, text
    ).paragraph_format.first_line_indent == Cm(0)


def test_long_first_line_is_not_a_heading_and_remains_in_body() -> None:
    long_line = "Д" * 121
    document = Document(build_post_docx(_post(text=f"{long_line}\n\nBody")))
    paragraphs = document.paragraphs

    assert sum(item.text == long_line for item in paragraphs) == 1
    long_paragraph = _paragraph(document, long_line)
    assert long_paragraph.runs[0].font.size == Pt(12.5)
    assert (
        long_paragraph.paragraph_format.first_line_indent.cm
        == pytest.approx(1.25, abs=0.001)
    )


def test_caption_is_used_as_heading_and_body_when_text_is_missing() -> None:
    document = Document(
        build_post_docx(_post(text=None, caption="Подпись\n\nПродолжение"))
    )
    paragraphs = [item.text for item in document.paragraphs]

    assert paragraphs.count("Подпись") == 1
    assert "Продолжение" in paragraphs


@pytest.mark.parametrize("photo_count", [1, 2, 3, 5])
def test_gallery_embeds_every_photo(photo_count: int) -> None:
    photos = [_photo(str(index)) for index in range(photo_count)]
    document = Document(build_post_docx(_post(), photos))

    assert len(document.inline_shapes) == photo_count
    assert len(document.tables) == 1
    assert len(document.tables[0].rows) == (photo_count + 1) // 2
    if photo_count % 2:
        assert len(document.tables[0].rows[-1].cells) == 2
        assert (
            document.tables[0].rows[-1].cells[0]._tc
            is document.tables[0].rows[-1].cells[1]._tc
        )


def test_gallery_inserts_photos_in_supplied_order() -> None:
    photos = [_photo("first", b"first"), _photo("second", b"second")]
    inserted = []

    def record_picture(_run, stream, **_kwargs):
        inserted.append(stream.read())

    with patch("docx.text.run.Run.add_picture", new=record_picture):
        build_post_docx(_post(), photos)

    assert inserted == [b"first", b"second"]


def test_photos_precede_body_in_document_xml() -> None:
    document = Document(
        build_post_docx(_post(text="Заголовок\n\nТекст"), [_photo()])
    )
    xml = document.element.body.xml

    assert xml.index("<w:tbl>") < xml.index("Текст")


def test_text_only_document_has_no_gallery() -> None:
    document = Document(build_post_docx(_post()))

    assert len(document.inline_shapes) == 0
    assert not document.tables


def test_photo_without_caption_has_no_fake_title_or_body() -> None:
    document = Document(
        build_post_docx(
            _post(text=None, caption=None, photo_count=1),
            [_photo()],
        )
    )
    paragraphs = [item.text for item in document.paragraphs]

    assert len(document.inline_shapes) == 1
    assert paragraphs == [
        "SOME CHANNEL",
        "07.09.2026 18:40 · Telegram",
        "https://t.me/somechannel/12345",
        "",
        FOOTER_TEXT,
    ]
