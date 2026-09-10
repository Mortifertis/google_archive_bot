"""Build in-memory DOCX files for forwarded Telegram posts."""

import re
from collections.abc import Sequence
from io import BytesIO

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.telegram_media import (TELEGRAM_DOWNLOAD_LIMIT_REASON, ArchivedMedia,
                                DownloadedMedia, DownloadedPhoto)
from app.telegram_parser import ForwardedPost

MAX_TITLE_LENGTH = 150
MAX_POST_HEADING_LENGTH = 120
FONT_NAME = "Noto Sans"
BODY_FONT_SIZE = Pt(12.5)


def _post_content(post: ForwardedPost) -> str:
    return post.text or post.caption or ""


def _first_meaningful_line(content: str) -> str:
    """Return the first non-empty line with whitespace normalized."""
    return next(
        (
            " ".join(line.split())
            for line in content.splitlines()
            if line.strip()
        ),
        "",
    )


def _post_heading(content: str) -> str:
    first_line = _first_meaningful_line(content)
    if len(first_line) <= MAX_POST_HEADING_LENGTH:
        return first_line
    return ""


def build_document_title(post: ForwardedPost) -> str:
    """Build a compact, single-line title from available post metadata."""
    date_part = (
        post.source_date.strftime("%Y-%m-%d")
        if post.source_date
        else "Telegram post"
    )
    parts = [date_part]
    if post.source_chat_title:
        parts.append(" ".join(post.source_chat_title.split()))

    first_line = _first_meaningful_line(_post_content(post))
    if first_line:
        parts.append(first_line)

    title = " — ".join(parts)
    if len(title) <= MAX_TITLE_LENGTH:
        return title
    return title[: MAX_TITLE_LENGTH - 1].rstrip() + "…"


class DocumentImageError(RuntimeError):
    """A photo could not be embedded in the archive document."""


def _set_run_font(run: object, size: Pt) -> None:
    run.font.name = FONT_NAME
    run.font.size = size
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT_NAME)


def _configure_document(document: DocumentType) -> None:
    section = document.sections[0]
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = BODY_FONT_SIZE
    fonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), FONT_NAME)


def _add_channel_heading(document: DocumentType, channel: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run(channel.upper())
    _set_run_font(run, Pt(9))
    run.bold = True
    run.font.color.rgb = RGBColor(68, 68, 68)


def _add_post_heading(document: DocumentType, heading: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(heading)
    _set_run_font(run, Pt(17))
    run.bold = True
    run.font.color.rgb = RGBColor(28, 28, 28)


def _add_metadata(document: DocumentType, post: ForwardedPost) -> None:
    published = (
        post.source_date.strftime("%d.%m.%Y %H:%M")
        if post.source_date
        else None
    )
    text = f"{published} · Telegram" if published else "Telegram"
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2 if post.source_url else 9)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run(text)
    _set_run_font(run, Pt(9))
    run.font.color.rgb = RGBColor(105, 105, 105)

    if post.source_url:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(9)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        run = paragraph.add_run(post.source_url)
        _set_run_font(run, Pt(9))
        run.font.color.rgb = RGBColor(58, 83, 112)


def _add_separator(document: DocumentType) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D9D9D9")
    borders.append(bottom)
    properties.append(borders)


def _remove_table_borders(table: object) -> None:
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is not None:
        properties.remove(borders)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    properties.append(borders)


def _configure_cell(cell: object) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge in ("top", "start", "bottom", "end"):
        margin = OxmlElement(f"w:{edge}")
        margin.set(qn("w:w"), "55")
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    properties.append(margins)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)


def format_duration(seconds: int | None) -> str | None:
    """Format a Telegram duration without external dependencies."""
    if seconds is None:
        return None
    hours, remainder = divmod(max(0, seconds), 3600)
    minutes, seconds = divmod(remainder, 60)
    if hours:
        return f"{hours}:{minutes:02d}:{seconds:02d}"
    return f"{minutes:02d}:{seconds:02d}"


def format_file_size(size: int | None) -> str | None:
    """Format byte size in compact decimal/binary-friendly units."""
    if size is None:
        return None
    if size < 1024 * 1024:
        return f"{size / 1000:.0f} KB"
    return f"{size / (1024 * 1024):.1f} MB"


def _add_hyperlink(paragraph: object, text: str, url: str) -> None:
    relationship = paragraph.part.relate_to(
        url,
        RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "3A5370")
    properties.append(color)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    properties.append(size)
    run.append(properties)
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _as_archived_media(item: object) -> ArchivedMedia:
    if isinstance(item, ArchivedMedia):
        return item
    if isinstance(item, DownloadedMedia):
        return ArchivedMedia(
            kind=item.kind,
            preview_data=(
                item.data if item.kind == "photo" else item.preview_data
            ),
            width=item.width,
            height=item.height,
            duration=item.duration,
            file_size=item.file_size,
            mime_type=item.mime_type,
            unavailable_reason=item.unavailable_reason,
        )
    photo = item
    return ArchivedMedia(
        kind="photo",
        preview_data=photo.data,
        width=photo.width,
        height=photo.height,
        duration=None,
        file_size=None,
    )


def _add_media_cell(cell: object, media: ArchivedMedia, width: object) -> None:
    _configure_cell(cell)
    paragraph = cell.paragraphs[0]
    if media.preview_data:
        run = paragraph.add_run()
        run.add_picture(BytesIO(media.preview_data), width=width)
    if media.kind == "photo":
        return

    label = "GIF" if media.mime_type == "image/gif" else {
        "video": "Видео",
        "animation": "Анимация",
    }[media.kind]
    details = [
        format_duration(media.duration),
        format_file_size(media.file_size),
    ]
    details = [value for value in details if value]
    line = f"▶ {label}"
    if details:
        line += " · " + " · ".join(details)
    metadata = cell.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_after = Pt(2)
    run = metadata.add_run(line)
    _set_run_font(run, Pt(9))
    run.bold = True
    run.font.color.rgb = RGBColor(68, 68, 68)
    if media.drive_web_url:
        link = cell.add_paragraph()
        link.alignment = WD_ALIGN_PARAGRAPH.CENTER
        noun = "видео" if media.kind == "video" else "анимацию"
        _add_hyperlink(
            link,
            f"Открыть {noun} в Google Drive",
            media.drive_web_url,
        )
    elif media.unavailable_reason == TELEGRAM_DOWNLOAD_LIMIT_REASON:
        warning = cell.add_paragraph()
        warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = warning.add_run(
            "Файл не архивирован: превышает лимит Telegram Bot API 20 MB."
        )
        _set_run_font(run, Pt(8))
        run.font.color.rgb = RGBColor(126, 74, 74)


def _add_gallery(
    document: DocumentType,
    items: Sequence[object],
) -> None:
    section = document.sections[0]
    available_width = (
        section.page_width - section.left_margin - section.right_margin
    )
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _remove_table_borders(table)
    column_width = (available_width - Cm(0.25)) // 2
    image_width = column_width - Cm(0.12)

    try:
        media_items = [_as_archived_media(item) for item in items]
        for index in range(0, len(media_items), 2):
            row = table.add_row()
            pair = media_items[index:index + 2]
            if len(pair) == 1:
                cell = row.cells[0].merge(row.cells[1])
                _add_media_cell(
                    cell,
                    pair[0],
                    available_width - Cm(0.15),
                )
                continue
            for cell, media in zip(row.cells, pair):
                cell.width = column_width
                _add_media_cell(cell, media, image_width)
    except Exception as error:
        raise DocumentImageError(
            "Could not embed a Telegram photo in DOCX"
        ) from error


def _body_paragraphs(content: str, heading: str) -> list[str]:
    if heading:
        lines = content.splitlines()
        first_index = next(
            (index for index, line in enumerate(lines) if line.strip()),
            None,
        )
        if first_index is not None:
            lines.pop(first_index)
        content = "\n".join(lines)
    return [
        block.strip()
        for block in re.split(r"\n\s*\n+", content.strip())
        if block.strip()
    ]


def _is_structural_paragraph(text: str) -> bool:
    return bool(
        re.match(r"^(?:[•●◦]|[-—*]\s|\d+[.)]\s)", text.lstrip())
    )


def _add_body(document: DocumentType, content: str, heading: str) -> None:
    for text in _body_paragraphs(content, heading):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.line_spacing = 1.2
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.first_line_indent = (
            Cm(0) if _is_structural_paragraph(text) else Cm(1.25)
        )
        run = paragraph.add_run(text)
        _set_run_font(run, BODY_FONT_SIZE)


def _add_footer(document: DocumentType) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.first_line_indent = Cm(0)
    run = paragraph.add_run("Сохранено через Telegram Archive Bot")
    _set_run_font(run, Pt(8))
    run.font.color.rgb = RGBColor(160, 160, 160)


def build_post_docx(
    post: ForwardedPost,
    photos: Sequence[DownloadedPhoto | DownloadedMedia | ArchivedMedia] = (),
) -> BytesIO:
    """Return a seekable DOCX stream containing post text and media."""
    document = Document()
    _configure_document(document)
    content = _post_content(post)
    heading = _post_heading(content)

    if post.source_chat_title:
        _add_channel_heading(document, post.source_chat_title)
    if heading:
        _add_post_heading(document, heading)
    _add_metadata(document, post)
    _add_separator(document)
    if photos:
        _add_gallery(document, photos)
    if content:
        _add_body(document, content, heading)
    _add_footer(document)

    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream
